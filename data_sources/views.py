from django.shortcuts import render

# Create your views here.
import openpyxl
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from enterprises.models import Enterprise, FinancialKPI, PropertyInfo, ProductInfo, GeoLocation
from .utils import parse_excel_file

@login_required
def generate_excel_template(request):
    if request.method != 'POST':
        return HttpResponse("Метод не разрешён", status=405)

    selected_models = request.POST.getlist('models')

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Словарь моделей (без Enterprise — он не нужен как отдельный лист для связи)
    model_map = {
        'financial_kpi': (FinancialKPI, "Финансовые показатели"),
        'property_info': (PropertyInfo, "Имущество"),
        'product_info': (ProductInfo, "Продукция"),
        'geo_location': (GeoLocation, "Геоданные"),
    }

    # Всегда добавляем ИНН как первый столбец
    for model_key in selected_models:
        if model_key == 'enterprise':
            # Лист "Основная информация" — включает ИНН как обычное поле
            ws = wb.create_sheet(title="Основная информация"[:31])
            headers = []
            for field in Enterprise._meta.fields:
                if field.name != 'id':
                    headers.append(field.verbose_name or field.name)
            ws.append(headers)
        elif model_key in model_map:
            # Все остальные листы — начинаются с ИНН
            model_class, sheet_title = model_map[model_key]
            ws = wb.create_sheet(title=sheet_title[:31])

            # Первый столбец — ИНН
            headers = ["ИНН"]

            # Остальные поля модели (без id и без enterprise ForeignKey)
            for field in model_class._meta.fields:
                if field.name not in ('id', 'enterprise'):
                    headers.append(field.verbose_name or field.name)

            ws.append(headers)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=industrial_template.xlsx'
    wb.save(response)
    return response

@login_required
def excel_upload_page(request):
    return render(request, 'admin/excel_upload.html')


@login_required
def upload_excel_file(request):
    if request.method == 'POST':
        file = request.FILES.get('file')
        if not file or not file.name.endswith('.xlsx'):
            return render(request, 'admin/excel_upload.html', {
                'report': {'errors': ['Пожалуйста, загрузите файл в формате .xlsx']}
            })

        report = parse_excel_file(file)
        return render(request, 'admin/excel_upload.html', {'report': report})

    return redirect('excel_upload_page')